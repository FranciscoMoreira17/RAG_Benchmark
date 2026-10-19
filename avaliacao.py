"""
avaliacao.py - Avaliação do Benchmark RAG (DRE)
===============================================
A unidade de recuperação é o DOCUMENTO, identificado pelo `guid`
(`doc_id` nos metadados). O alvo de cada pergunta está em
`documentos_esperados` (lista de guids).

Camadas de avaliação:
  1. RETRIEVAL (determinístico) - precision@k, recall@k, MRR@k, nDCG@k,
     hit@k, com IC95% por bootstrap. Estratificado por tipo e camada.
  2. GERAÇÃO - juiz LLM (juiz_llm.py), só perguntas para_geracao=True.
  3. CLÁSSICAS - BLEU, ROUGE-L, BERTScore (anexo).

SUBCONJUNTOS (campo `subconjunto` do resultado ou do dataset):
  - "original"   (100 perguntas): tabelas 1-4 e curva hit@k.
  - "articulado" (30 perguntas):  reportado à parte (tabela 8).
  - "total":                      só informativo (tabelas 8 e 9).

TESTES (tabela 9), só para os pares definidos a priori em PARES_POR_ESTUDO:
  - McNemar exato (binomial) sobre o hit@5 por pergunta.
  - Teste de permutação emparelhado (sign-flip) sobre MRR@5 e nDCG@5.
  Correção de Holm dentro de cada estudo x subconjunto. O estudo é o
  nome da pasta de resultados (ex.: resultados/Top100/Estudo1).
  Com várias repetições de uma condição: hit por voto da maioria,
  MRR/nDCG pela média entre repetições.

Uso:
    python avaliacao.py --pasta resultados/Top100/Estudo1 --so-retrieval
    python avaliacao.py --ficheiro resultados/Top100/Estudo1/X.json
    python avaliacao.py --comparar --pasta resultados/Top100/Estudo1
    python avaliacao.py --ficheiro X.json --langsmith
"""

import os
import sys
import json
import glob
import math
import random
import argparse
import statistics
from collections import defaultdict

import pandas as pd
from dotenv import load_dotenv

_DIR = os.path.dirname(os.path.abspath(__file__))
if _DIR not in sys.path:
    sys.path.insert(0, _DIR)

load_dotenv()

# =====================================================================
# Configuração
# =====================================================================
RESULTADOS_DIR = os.getenv("RESULTADOS_DIR",
                           os.path.join(_DIR, "resultados", "Top100", "Estudo3"))
DATASET_PATH = os.getenv("DATASET_PATH", os.path.join(_DIR, "dataset_dre_manual.json"))
JUDGE_MODEL_GROQ = os.getenv("JUDGE_MODEL_GROQ", "openai/gpt-oss-120b")

K = 5                              
CURVA_K = (1, 3, 5, 10, 20)
SUBCONJUNTOS = ("original", "articulado")
ALPHA = 0.05
N_BOOTSTRAP = 2000
N_PERMUTACOES = 10000
SEED = 42


PARES_POR_ESTUDO = {
    "Estudo1": [
        # Isolamento de fatores (referência: ULS-Qwen-híbrido)
        ("UnitLevelSegmentationQwen-hibrido", "overlap-qwen-hibrido"),             # segmentação
        ("UnitLevelSegmentationQwen-hibrido", "UnitLevelSegmentationQwen-denso"),  # híbrido vs denso
        ("UnitLevelSegmentationQwen-hibrido", "UnitLevelSegmentation-esparso"),    # híbrido vs esparso
        ("UnitLevelSegmentationQwen-hibrido", "UnitLevelSegmentation-hibrido"),    # embedder
        # Baseline
        ("UnitLevelSegmentationQwen-hibrido", "langchain-qwen"),
    ],
    "Estudo2": [
        ("UnitLevelSegmentationQwen512-denso", "UnitLevelSegmentation-denso"),     # embedder (janela igual)
        ("UnitLevelSegmentationQwen-denso", "UnitLevelSegmentationQwen512-denso"), # janela
    ],
    "Estudo3": [
        ("UnitLevelSegmentationQwen-hibrido", "documento-qwen-hibrido"),
        ("ULSv2Qwen-hibrido", "UnitLevelSegmentationQwen-hibrido"),
        ("ULSv2Qwen-hibrido", "documento-qwen-hibrido"),
        ("estrutural-qwen-enriquecido", "UnitLevelSegmentationQwen-hibrido"),
    ],
}


def _get_client_groq():
    """Cliente Groq (OpenAI-compatible) usado pelo juiz_llm."""
    from openai import OpenAI
    return OpenAI(base_url="https://api.groq.com/openai/v1",
                  api_key=os.environ["GROQ_API_KEY"])


# =====================================================================
# Estatística
# =====================================================================
def bootstrap_ic(valores: list[float], n: int = N_BOOTSTRAP,
                 alpha: float = ALPHA) -> list[float]:
    """IC (percentil) da média por bootstrap."""
    if not valores:
        return [0.0, 0.0]
    if len(valores) == 1:
        return [round(valores[0], 4)] * 2
    rng = random.Random(SEED)
    m = len(valores)
    medias = sorted(sum(valores[rng.randrange(m)] for _ in range(m)) / m
                    for _ in range(n))
    return [round(medias[int((alpha / 2) * n)], 4),
            round(medias[int((1 - alpha / 2) * n) - 1], 4)]


def mcnemar_exato(b: int, c: int) -> float:
    """p bilateral do McNemar exato (binomial, p=0.5) sobre os discordantes."""
    n = b + c
    if n == 0:
        return 1.0
    cauda = sum(math.comb(n, i) for i in range(min(b, c) + 1)) / (2 ** n)
    return min(1.0, 2 * cauda)


def permutacao_emparelhada(a: list[float], b: list[float],
                           n: int = N_PERMUTACOES) -> float:
    """p bilateral do teste de permutação emparelhado (sign-flip) sobre a média das diferenças."""
    difs = [x - y for x, y in zip(a, b)]
    if not difs or all(d == 0 for d in difs):
        return 1.0
    obs = abs(statistics.mean(difs))
    rng = random.Random(SEED)
    extremos = sum(
        1 for _ in range(n)
        if abs(statistics.mean(d if rng.random() < 0.5 else -d for d in difs)) >= obs - 1e-12
    )
    return (extremos + 1) / (n + 1)


def holm(pvalores: list[float]) -> list[float]:
    """Correção de Holm-Bonferroni (p ajustados na ordem original)."""
    m = len(pvalores)
    ordem = sorted(range(m), key=lambda i: pvalores[i])
    ajust, maximo = [0.0] * m, 0.0
    for pos, i in enumerate(ordem):
        maximo = max(maximo, min(1.0, (m - pos) * pvalores[i]))
        ajust[i] = maximo
    return ajust


# =====================================================================
# 1. RETRIEVAL
# =====================================================================
def _docs_distintos(resultado: dict) -> list[str]:
    """guids distintos, pela ordem do ranking (vários segmentos do mesmo doc contam uma vez)."""
    vistos, docs = set(), []
    for m in resultado.get("metadados", []):
        g = m.get("doc_id", "")
        if g and g not in vistos:
            vistos.add(g)
            docs.append(g)
    return docs


def _dcg(rels: list[float]) -> float:
    return sum(r / math.log2(i + 2) for i, r in enumerate(rels))


def metricas_pergunta(resultado: dict, k: int) -> dict | None:
    """Métricas de uma pergunta. None se não tiver documentos_esperados."""
    alvos = set(resultado.get("documentos_esperados") or [])
    if not alvos:
        return None
    docs = _docs_distintos(resultado)
    topk = docs[:k]
    rels = [1.0 if g in alvos else 0.0 for g in topk]
    acertos = len(alvos & set(topk))
    ideal = _dcg([1.0] * min(len(alvos), k))
    return {
        "precision": sum(rels) / k,
        "recall": acertos / len(alvos),
        "mrr": next((1 / (i + 1) for i, x in enumerate(rels) if x > 0), 0.0),
        "ndcg": _dcg(rels) / ideal if ideal > 0 else 0.0,
        "hit": 1.0 if acertos else 0.0,
        "ndist": len(docs),
    }


def metricas_retrieval(resultados: list[dict], k: int = K) -> dict:
    """
    Agregado por condição. `docs_distintos_min` é o mínimo, entre perguntas,
    do número de documentos distintos em TODOS os segmentos recuperados
    (diagnóstico de concentração do ranking num número reduzido de documentos).
    """
    por_q = [m for m in (metricas_pergunta(r, k) for r in resultados) if m]
    if not por_q:
        return {}
    col = lambda c: [m[c] for m in por_q]
    return {
        f"precision@{k}": statistics.mean(col("precision")),
        f"recall@{k}": statistics.mean(col("recall")),
        "mrr": statistics.mean(col("mrr")),
        f"ndcg@{k}": statistics.mean(col("ndcg")),
        f"hit_rate@{k}": statistics.mean(col("hit")),
        "n_queries": len(por_q),
        "docs_distintos_min": min(col("ndist")),
        "ic95_recall": bootstrap_ic(col("recall")),
        "ic95_hit": bootstrap_ic(col("hit")),
    }


def _rotulo_camada(valor) -> str:
    return {1: "facil", 2: "dificil"}.get(valor, f"camada{valor}")


def estratificar(resultados: list[dict], chave, k: int = K) -> dict:
    """Agrupa por `chave(r) -> rótulo` e calcula o retrieval de cada grupo."""
    grupos = defaultdict(list)
    for r in resultados:
        grupos[chave(r)].append(r)
    out = {}
    for rotulo in sorted(grupos, key=str):
        m = metricas_retrieval(grupos[rotulo], k)
        if m:
            out[rotulo] = m
    return out


_por_tipo = lambda r: f"tipo={r.get('tipo', '?')}"
_por_camada = lambda r: f"camada={_rotulo_camada(r.get('camada'))}"
_por_tipo_camada = lambda r: f"tipo={r.get('tipo', '?')}|camada={_rotulo_camada(r.get('camada'))}"


# =====================================================================
# Subconjuntos
# =====================================================================
_MAPA_SUB = None


def _mapa_subconjunto() -> dict:
    """id -> subconjunto, lido do dataset (uma vez)."""
    global _MAPA_SUB
    if _MAPA_SUB is None:
        _MAPA_SUB = {}
        try:
            with open(DATASET_PATH, "r", encoding="utf-8") as f:
                for q in json.load(f):
                    if q.get("subconjunto"):
                        _MAPA_SUB[str(q["id"])] = q["subconjunto"]
        except FileNotFoundError:
            print(f"    [AVISO] Dataset não encontrado ({DATASET_PATH}); "
                  f"subconjunto só a partir dos resultados.")
    return _MAPA_SUB


def subconjunto_de(r: dict) -> str:
    return (r.get("subconjunto")
            or _mapa_subconjunto().get(str(r.get("id")))
            or "original")


def verificar_orcamento(resultados: list[dict]) -> dict:
    toks = [r["tokens_contexto"] for r in resultados if r.get("tokens_contexto")]
    blocos = [r["n_blocos_usados"] for r in resultados if r.get("n_blocos_usados")]
    if not toks:
        return {}
    return {
        "tokens_media": statistics.mean(toks),
        "tokens_desvio": statistics.pstdev(toks) if len(toks) > 1 else 0.0,
        "tokens_min": min(toks),
        "tokens_max": max(toks),
        "blocos_media": statistics.mean(blocos) if blocos else 0,
    }

import re
_MAPA_ART = None

def artigo_esperado_de(r: dict) -> str:
    global _MAPA_ART
    if r.get("artigo_esperado"):
        return r["artigo_esperado"]
    if _MAPA_ART is None:
        try:
            with open(DATASET_PATH, "r", encoding="utf-8") as f:
                _MAPA_ART = {str(q["id"]): q.get("artigo_esperado", "") for q in json.load(f)}
        except FileNotFoundError:
            _MAPA_ART = {}
    return _MAPA_ART.get(str(r.get("id")), "")


def artigo_no_contexto(r: dict):
    """1 se o cabeçalho do artigo esperado está num bloco do documento certo entregue ao gerador."""
    m = re.search(r"\d+", artigo_esperado_de(r) or "")
    if not m:
        return None
    padrao = re.compile(rf"(?m)^\s*#{{0,6}}\s*\**\s*Art(?:igo)?\.?\s*{m.group()}\s*[\.\-]?\s*[º°oO]",
                        re.IGNORECASE)
    alvos = set(r.get("documentos_esperados") or [])
    for txt, meta in zip(r.get("contextos") or [], r.get("metadados") or []):
        if meta.get("doc_id") in alvos and padrao.search(txt or ""):
            return 1
    return 0

# =====================================================================
# 2-3. GERAÇÃO e CLÁSSICAS
# =====================================================================
def _sem_erro(resultados: list[dict]) -> list[dict]:
    return [r for r in resultados if not r.get("resposta", "").startswith("ERRO")]


def _validos_geracao(resultados: list[dict]) -> list[dict]:
    """Perguntas para geração, sem erro e com ground truth."""
    return [r for r in _sem_erro(resultados)
            if r.get("para_geracao") and r.get("ground_truth")]


def avaliar_classicas(resultados: list[dict]) -> dict:
    from rouge_score import rouge_scorer
    from nltk.translate.bleu_score import sentence_bleu, SmoothingFunction

    validos = _validos_geracao(resultados)
    if not validos:
        return {}

    scorer = rouge_scorer.RougeScorer(["rougeL"], use_stemmer=True)
    smooth = SmoothingFunction().method1
    bleus = [sentence_bleu([r["ground_truth"].split()], r["resposta"].split(),
                           smoothing_function=smooth) for r in validos]
    rouges = [scorer.score(r["ground_truth"], r["resposta"])["rougeL"].fmeasure
              for r in validos]

    out = {
        "bleu": statistics.mean(bleus),
        "rouge_l": statistics.mean(rouges),
        "rouge_l_ic95": bootstrap_ic(rouges),
        "_nota": "Sobreposição lexical com referência única. Anexo, não métrica principal.",
    }
    try:
        from bert_score import score as bert_fn
        print("    [Clássicas] BERTScore...")
        _, _, f1 = bert_fn([r["resposta"] for r in validos],
                           [r["ground_truth"] for r in validos],
                           lang="pt", verbose=False)
        out["bertscore_f1"] = float(f1.mean())
    except Exception as e:
        print(f"    [Clássicas] BERTScore indisponível: {e}")
    return out


# =====================================================================
# LangSmith (opcional, --langsmith)
# =====================================================================
def avaliar_com_langsmith(caminho: str):
    try:
        from langsmith import Client
        from langsmith.evaluation import evaluate
    except ImportError:
        print("    [LangSmith] Não instalado — skip")
        return
    if not (os.getenv("LANGSMITH_API_KEY") or os.getenv("LANGCHAIN_API_KEY")):
        print("    [LangSmith] Sem API key — skip")
        return

    with open(caminho, "r", encoding="utf-8") as f:
        dados = json.load(f)
    condicao = dados.get("condicao", dados["framework"])
    resultados = dados["resultados"]
    client = Client()

    nome_ds = os.getenv("LANGSMITH_DATASET", "rag-benchmark-dre")
    if client.has_dataset(dataset_name=nome_ds):
        client.read_dataset(dataset_name=nome_ds)
    else:
        ds = client.create_dataset(dataset_name=nome_ds)
        client.create_examples(
            dataset_id=ds.id,
            examples=[{
                "inputs": {"query": r["query"]},
                "outputs": {"ground_truth": r.get("ground_truth", ""),
                            "documentos_esperados": r.get("documentos_esperados", [])},
                "metadata": {"tipo": r["tipo"], "id": r["id"], "camada": r.get("camada")},
            } for r in resultados if r.get("documentos_esperados")],
        )
        print(f"    [LangSmith] Dataset criado: {nome_ds}")

    por_query = {r["query"]: r for r in resultados}

    def target(inputs: dict) -> dict:
        r = por_query.get(inputs["query"], {})
        return {"resposta": r.get("resposta", ""),
                "metadados": r.get("metadados", []),
                "tokens_contexto": r.get("tokens_contexto", 0)}

    def hit_guid(outputs, reference_outputs):
        alvos = set(reference_outputs.get("documentos_esperados") or [])
        if not alvos:
            return {"key": "hit_guid", "score": None}
        obtidos = {m.get("doc_id") for m in outputs.get("metadados", [])}
        return {"key": "hit_guid", "score": 1.0 if alvos & obtidos else 0.0}

    def recall_guid(outputs, reference_outputs):
        alvos = set(reference_outputs.get("documentos_esperados") or [])
        if not alvos:
            return {"key": "recall_guid", "score": None}
        obtidos = {m.get("doc_id") for m in outputs.get("metadados", [])}
        return {"key": "recall_guid", "score": len(alvos & obtidos) / len(alvos)}

    def sem_erro(outputs):
        return {"key": "sem_erro",
                "score": 0.0 if outputs["resposta"].startswith("ERRO") else 1.0}

    evaluate(target, data=nome_ds,
             evaluators=[hit_guid, recall_guid, sem_erro],
             experiment_prefix=condicao,
             metadata={"condicao": condicao, **dados.get("config", {})},
             max_concurrency=4)
    print(f"    [LangSmith] Experiência registada: {condicao}")


# =====================================================================
# Avaliação de um ficheiro
# =====================================================================
def avaliar_ficheiro(caminho: str, so_retrieval: bool = False,
                     usar_langsmith: bool = False) -> dict:
    print(f"\n{'=' * 62}\n  AVALIAÇÃO: {os.path.basename(caminho)}\n{'=' * 62}")

    with open(caminho, "r", encoding="utf-8") as f:
        dados = json.load(f)
    resultados = dados["resultados"]
    validos = _sem_erro(resultados)

    def media_tempo(campo):
        vals = [r[campo] for r in validos if r.get(campo) is not None]
        return statistics.mean(vals) if vals else 0

    por_sub = defaultdict(list)
    for r in resultados:
        por_sub[subconjunto_de(r)].append(r)
    originais = por_sub.get("original", [])

    resumo = {
        "condicao": dados.get("condicao", dados["framework"]),
        "framework": dados["framework"],
        "repeticao": dados.get("repeticao", 1),
        "config": dados.get("config", {}),
        "n_queries": len(resultados),
        "n_erros": len(resultados) - len(validos),
        "n_por_subconjunto": {s: len(v) for s, v in por_sub.items()},
        "orcamento_contexto": verificar_orcamento(resultados),
        "tempos": {
            "retrieval_medio_s": media_tempo("tempo_retrieval_s"),
            "geracao_media_s": media_tempo("tempo_geracao_s"),
            "total_medio_s": media_tempo("tempo_total_s"),
        },
    }

    # [1] Retrieval: tabelas principais só sobre o subconjunto original
    print(f"\n  [1] Retrieval — {resumo['n_por_subconjunto']}")
    resumo["retrieval"] = metricas_retrieval(originais)
    resumo["retrieval_estratificado"] = estratificar(originais, _por_tipo)
    resumo["retrieval_por_camada"] = estratificar(originais, _por_camada)
    resumo["retrieval_tipo_camada"] = estratificar(originais, _por_tipo_camada)
    resumo["retrieval_curva"] = {str(kk): metricas_retrieval(originais, kk) for kk in CURVA_K}
    resumo["orcamento_por_subconjunto"] = {s: verificar_orcamento(rs) for s, rs in por_sub.items()}
    
    # Tabela 8: curva por subconjunto (+ total)
    grupos = [(s, por_sub[s]) for s in SUBCONJUNTOS if por_sub.get(s)]
    if len(grupos) > 1:
        grupos.append(("total", resultados))
    resumo["retrieval_por_subconjunto"] = {
        s: {str(kk): metricas_retrieval(rs, kk) for kk in CURVA_K} for s, rs in grupos
    }

    # Tabela 9: métricas por pergunta (hit, mrr, ndcg @K) + subconjunto de cada id
    resumo["por_pergunta"] = {
        str(r.get("id")): {c: m[c] for c in ("hit", "mrr", "ndcg")}
        for r in resultados if (m := metricas_pergunta(r, K))
    }
    resumo["subconjunto_por_id"] = {str(r.get("id")): subconjunto_de(r) for r in resultados}

    if not so_retrieval:
        from juiz_llm import avaliar_llm_juiz
        print("\n  [2] Juiz LLM (faithfulness, relevancy, correctness)...")

        resumo["juiz_por_subconjunto"] = {}
        for s in SUBCONJUNTOS:
            rs = por_sub.get(s, [])
            if _validos_geracao(rs):
                print(f"    [Juiz] subconjunto={s}: {len(_validos_geracao(rs))} perguntas")
                j = avaliar_llm_juiz(rs, _get_client_groq)
                if j:
                    resumo["juiz_por_subconjunto"][s] = j
        resumo["juiz_llm"] = resumo["juiz_por_subconjunto"].get("original", {})

        art = [r for r in por_sub.get("articulado", []) if r.get("para_geracao")]
        vals = [v for v in (artigo_no_contexto(r) for r in art) if v is not None]
        if vals:
            resumo["artigo_no_contexto"] = {"taxa": statistics.mean(vals), "n": len(vals)}

        print("\n  [3] Clássicas (BLEU/ROUGE/BERTScore, anexo)...")
        resumo["classicas"] = avaliar_classicas(originais)

    if usar_langsmith:
        print("\n  [4] LangSmith...")
        avaliar_com_langsmith(caminho)

    destino = caminho.replace(".json", "_avaliacao.json")
    with open(destino, "w", encoding="utf-8") as f:
        json.dump(resumo, f, ensure_ascii=False, indent=2)

    imprimir_resumo(resumo)
    print(f"\n  Guardado: {destino}")
    return resumo


def imprimir_resumo(resumo: dict):
    print(f"\n  Condição: {resumo['condicao']}  "
          f"({resumo['n_queries']} queries, {resumo['n_erros']} erros)")

    orc = resumo.get("orcamento_contexto") or {}
    if orc:
        print(f"  Contexto: {orc['tokens_media']:.0f} tokens médios "
              f"(σ={orc['tokens_desvio']:.0f})")

    m = resumo.get("retrieval") or {}
    if m:
        print(f"\n  Retrieval (original, n={m['n_queries']}):")
        for kk, vv in m.items():
            if kk != "n_queries" and not kk.startswith("ic95"):
                print(f"    {kk:<20} {vv:.4f}")
        print(f"    hit@{K} IC95%         {m.get('ic95_hit')}")

    for chave, mm in (resumo.get("retrieval_estratificado") or {}).items():
        print(f"\n  Retrieval [{chave}]  n={mm['n_queries']}  "
              f"hit@{K}={mm[f'hit_rate@{K}']:.3f}  recall@{K}={mm[f'recall@{K}']:.3f}  "
              f"ndcg@{K}={mm[f'ndcg@{K}']:.3f}")

    for bloco in ("juiz_llm", "classicas"):
        for kk, vv in (resumo.get(bloco) or {}).items():
            if not kk.startswith("_") and not kk.endswith("_ic95") and isinstance(vv, (int, float)):
                print(f"    [{bloco}] {kk:<32} {vv:.4f}")


# =====================================================================
# Comparação (Excel com 9 tabelas empilhadas)
# =====================================================================
def comparar(pasta: str = None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    pasta = pasta or RESULTADOS_DIR
    estudo = os.path.basename(os.path.normpath(pasta))
    ficheiros = glob.glob(os.path.join(pasta, "*_avaliacao.json"))
    if not ficheiros:
        print("Sem avaliações. Corre primeiro --pasta <pasta do estudo>")
        return

    por_condicao = defaultdict(list)
    for p in sorted(ficheiros):
        with open(p, "r", encoding="utf-8") as f:
            r = json.load(f)
        por_condicao[r["condicao"]].append(r)

    # ---------------- utilitários ----------------
    def obter(d, *caminho):
        for p in caminho:
            d = d.get(p) if isinstance(d, dict) else None
        return d

    def media_desvio(reps, getter):
        vals = [v for v in (getter(r) for r in reps) if v is not None]
        if not vals:
            return None, None
        return statistics.mean(vals), (statistics.pstdev(vals) if len(vals) > 1 else 0.0)

    def media(reps, getter):
        return media_desvio(reps, getter)[0]

    def g_juiz(r, m, sub):
        d = obter(r, "juiz_llm", "_detalhe") or {}
        if m in d:
            return d[m].get(sub)
        return obter(r, "juiz_llm", m) if sub == "media" else None

    # ---------------- tabelas ----------------
    def tab_geral():
        L = []
        for cond, reps in sorted(por_condicao.items()):
            hit_m, hit_s = media_desvio(reps, lambda r: obter(r, "retrieval", f"hit_rate@{K}"))
            L.append({
                "condicao": cond, "n_exec": len(reps),
                "n_queries": obter(reps[0], "retrieval", "n_queries"),
                f"hit@{K}": hit_m, f"hit@{K}_ic95": obter(reps[0], "retrieval", "ic95_hit"),
                f"hit@{K}_sd_exec": hit_s,
                f"recall@{K}": media(reps, lambda r: obter(r, "retrieval", f"recall@{K}")),
                "mrr": media(reps, lambda r: obter(r, "retrieval", "mrr")),
                f"ndcg@{K}": media(reps, lambda r: obter(r, "retrieval", f"ndcg@{K}")),
                f"precision@{K}": media(reps, lambda r: obter(r, "retrieval", f"precision@{K}")),
            })
        return pd.DataFrame(L)

    def tab_curva():
        L = []
        for cond, reps in sorted(por_condicao.items()):
            linha = {"condicao": cond}
            for kk in CURVA_K:
                linha[f"hit@{kk}"] = media(
                    reps, lambda r, kk=kk: obter(r, "retrieval_curva", str(kk), f"hit_rate@{kk}"))
            linha["docs_distintos_min"] = media(reps, lambda r: obter(r, "retrieval", "docs_distintos_min"))
            L.append(linha)
        return pd.DataFrame(L)

    def tab_estrato(bloco, ordem, etiqueta):
        L = []
        for cond, reps in sorted(por_condicao.items()):
            chaves = set()
            for r in reps:
                chaves.update((r.get(bloco) or {}).keys())
            seq = [c for c in ordem if c in chaves] + sorted(c for c in chaves if c not in ordem)
            for chave in seq:
                g = lambda r, m: obter(r, bloco, chave, m)
                L.append({
                    "condicao": cond,
                    etiqueta: chave.split("=", 1)[-1],
                    "n_queries": g(reps[0], "n_queries"),
                    f"hit@{K}": media(reps, lambda r: g(r, f"hit_rate@{K}")),
                    f"recall@{K}": media(reps, lambda r: g(r, f"recall@{K}")),
                    "mrr": media(reps, lambda r: g(r, "mrr")),
                    f"ndcg@{K}": media(reps, lambda r: g(r, f"ndcg@{K}")),
                })
        return pd.DataFrame(L)

    def tab_geracao():
        L = []
        for cond, reps in sorted(por_condicao.items()):
            linha = {"condicao": cond, "n_geracao": g_juiz(reps[0], "faithfulness", "n")}
            for m, abrev in (("faithfulness", "faith"), ("answer_relevancy", "relev"),
                             ("correctness", "correct")):
                linha[m] = media(reps, lambda r, m=m: g_juiz(r, m, "media"))
                linha[f"{abrev}_sd"] = media(reps, lambda r, m=m: g_juiz(r, m, "desvio"))
            L.append(linha)
        return pd.DataFrame(L)

    def tab_classicas():
        return pd.DataFrame([{
            "condicao": cond,
            **{m: media(reps, lambda r, m=m: obter(r, "classicas", m))
               for m in ("bleu", "rouge_l", "bertscore_f1")},
        } for cond, reps in sorted(por_condicao.items())])

    def tab_geracao_sub():
        L = []
        for cond, reps in sorted(por_condicao.items()):
            for s in SUBCONJUNTOS:
                def gj(r, m, sub, s=s):
                    d = obter(r, "juiz_por_subconjunto", s, "_detalhe") or {}
                    if m in d:
                        return d[m].get(sub)
                    return obter(r, "juiz_por_subconjunto", s, m) if sub == "media" else None
                if gj(reps[0], "faithfulness", "media") is None:
                    continue
                linha = {"condicao": cond, "subconjunto": s,
                         "n_geracao": gj(reps[0], "faithfulness", "n")}
                for m in ("faithfulness", "answer_relevancy", "correctness"):
                    linha[m] = media(reps, lambda r, m=m: gj(r, m, "media"))
                linha["artigo_no_contexto"] = (media(reps, lambda r: obter(r, "artigo_no_contexto", "taxa"))
                                               if s == "articulado" else None)
                linha["tokens_ctx"] = media(reps, lambda r: obter(r, "orcamento_por_subconjunto", s, "tokens_media"))
                linha["blocos_usados"] = media(reps, lambda r: obter(r, "orcamento_por_subconjunto", s, "blocos_media"))
                L.append(linha)
        return pd.DataFrame(L)

    def tab_custo():
        return pd.DataFrame([{
            "condicao": cond,
            "n_queries": media(reps, lambda r: r.get("n_queries")),
            "n_erros": media(reps, lambda r: r.get("n_erros")),
            "tokens_ctx": media(reps, lambda r: obter(r, "orcamento_contexto", "tokens_media")),
            "lat_retrieval_s": media(reps, lambda r: obter(r, "tempos", "retrieval_medio_s")),
            "lat_geracao_s": media(reps, lambda r: obter(r, "tempos", "geracao_media_s")),
            "lat_total_s": media(reps, lambda r: obter(r, "tempos", "total_medio_s")),
        } for cond, reps in sorted(por_condicao.items())])

    def tab_subconjunto():
        L = []
        for cond, reps in sorted(por_condicao.items()):
            subs = set()
            for r in reps:
                subs.update((r.get("retrieval_por_subconjunto") or {}).keys())
            for sub in [s for s in (*SUBCONJUNTOS, "total") if s in subs]:
                g = lambda r, kk, m: obter(r, "retrieval_por_subconjunto", sub, str(kk), m)
                linha = {"condicao": cond, "subconjunto": sub,
                         "n_queries": g(reps[0], K, "n_queries")}
                for kk in CURVA_K:
                    linha[f"hit@{kk}"] = media(reps, lambda r, kk=kk: g(r, kk, f"hit_rate@{kk}"))
                linha["mrr"] = media(reps, lambda r: g(r, K, "mrr"))
                linha[f"ndcg@{K}"] = media(reps, lambda r: g(r, K, f"ndcg@{K}"))
                linha["docs_distintos_min"] = media(reps, lambda r: g(r, K, "docs_distintos_min"))
                L.append(linha)
        return pd.DataFrame(L)

    def por_pergunta_agregado(cond):
        """id -> {hit (voto da maioria), mrr, ndcg (médias)} entre repetições."""
        acum = defaultdict(lambda: defaultdict(list))
        for r in por_condicao.get(cond, []):
            for qid, m in (r.get("por_pergunta") or {}).items():
                for c, v in m.items():
                    acum[qid][c].append(v)
        return {qid: {"hit": 1 if statistics.mean(d["hit"]) >= 0.5 else 0,
                      "mrr": statistics.mean(d["mrr"]),
                      "ndcg": statistics.mean(d["ndcg"])}
                for qid, d in acum.items()}

    def tab_testes():
        pares = PARES_POR_ESTUDO.get(estudo)
        if not pares:
            print(f"    [Testes] sem pares definidos para '{estudo}'")
            return pd.DataFrame()

        # Validação única (avisa uma vez por par, não uma vez por subconjunto)
        validos = []
        for a, b in pares:
            em_falta = [c for c in (a, b) if c not in por_condicao]
            if em_falta:
                print(f"    [Testes] par ignorado em {estudo}, condição em falta: {', '.join(em_falta)}")
            else:
                validos.append((a, b))

        sub_por_id = {}
        for reps in por_condicao.values():
            for r in reps:
                sub_por_id.update(r.get("subconjunto_por_id") or {})
        dados = {c: por_pergunta_agregado(c) for par in validos for c in par}

        L = []
        for sub in (*SUBCONJUNTOS, "total"):
            linhas = []
            for a, b in validos:
                pa, pb = dados[a], dados[b]
                ids = sorted(q for q in pa if q in pb
                             and (sub == "total" or sub_por_id.get(q, "original") == sub))
                if not ids:
                    continue
                b_ = sum(1 for q in ids if pa[q]["hit"] and not pb[q]["hit"])
                c_ = sum(1 for q in ids if pb[q]["hit"] and not pa[q]["hit"])
                mrr_a, mrr_b = [pa[q]["mrr"] for q in ids], [pb[q]["mrr"] for q in ids]
                ndcg_a, ndcg_b = [pa[q]["ndcg"] for q in ids], [pb[q]["ndcg"] for q in ids]
                linhas.append({
                    "subconjunto": sub, "condicao_A": a, "condicao_B": b, "n": len(ids),
                    f"hit@{K}_A": sum(pa[q]["hit"] for q in ids) / len(ids),
                    f"hit@{K}_B": sum(pb[q]["hit"] for q in ids) / len(ids),
                    "A_certo_B_errado": b_, "B_certo_A_errado": c_,
                    "p_mcnemar": mcnemar_exato(b_, c_),
                    "delta_mrr": statistics.mean(mrr_a) - statistics.mean(mrr_b),
                    "p_perm_mrr": permutacao_emparelhada(mrr_a, mrr_b),
                    f"delta_ndcg@{K}": statistics.mean(ndcg_a) - statistics.mean(ndcg_b),
                    f"p_perm_ndcg@{K}": permutacao_emparelhada(ndcg_a, ndcg_b),
                })
            if not linhas:
                continue
            for col_p, col_h in (("p_mcnemar", "p_holm_mcnemar"),
                                 ("p_perm_mrr", "p_holm_mrr"),
                                 (f"p_perm_ndcg@{K}", f"p_holm_ndcg@{K}")):
                for x, pa_ in zip(linhas, holm([x[col_p] for x in linhas])):
                    x[col_h] = pa_
            for x in linhas:
                x[f"signif_hit_{ALPHA}"] = "sim" if x["p_holm_mcnemar"] < ALPHA else "não"
            L.extend(linhas)
        return pd.DataFrame(L)

    tabelas = [
        ("1. RETRIEVAL - GERAL (subconjunto original)", tab_geral()),
        ("1b. RETRIEVAL - CURVA hit@k (subconjunto original)", tab_curva()),
        ("2. RETRIEVAL - POR TIPO",
         tab_estrato("retrieval_estratificado",
                     ["tipo=factual/lookup", "tipo=semantica"], "tipo")),
        ("3. RETRIEVAL - POR CAMADA (facil vs dificil)",
         tab_estrato("retrieval_por_camada", ["camada=facil", "camada=dificil"], "camada")),
        ("4. RETRIEVAL - TIPO x CAMADA",
         tab_estrato("retrieval_tipo_camada",
                     ["tipo=factual/lookup|camada=facil", "tipo=factual/lookup|camada=dificil",
                      "tipo=semantica|camada=facil", "tipo=semantica|camada=dificil"],
                     "subcategoria")),
        ("5. GERACAO - JUIZ LLM", tab_geracao()),
        ("5b. GERACAO - JUIZ LLM POR SUBCONJUNTO (+ artigo no contexto)", tab_geracao_sub()),
        ("6. METRICAS CLASSICAS (anexo)", tab_classicas()),
        ("7. CUSTO / EXECUCAO", tab_custo()),
        ("8. RETRIEVAL - POR SUBCONJUNTO (original / articulado / total)", tab_subconjunto()),
        (f"9. TESTES EMPARELHADOS - {estudo} (McNemar hit@{K}; permutação MRR e nDCG@{K}; "
         f"Holm por subconjunto)", tab_testes()),
    ]

    # ---------------- Excel ----------------
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "comparacao"
    titulo_font = Font(name="Arial", bold=True, size=12, color="1F4E78")
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    body_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="D9D9D9")
    borda = Border(left=thin, right=thin, top=thin, bottom=thin)

    linha = 1
    larguras = defaultdict(int)
    for titulo, df in tabelas:
        ws.cell(row=linha, column=1, value=titulo).font = titulo_font
        linha += 1
        if df is None or df.empty:
            ws.cell(row=linha, column=1, value="(sem dados)").font = body_font
            linha += 3
            continue
        for j, col in enumerate(df.columns, start=1):
            c = ws.cell(row=linha, column=j, value=str(col))
            c.font, c.fill, c.border = header_font, header_fill, borda
            c.alignment = Alignment(horizontal="center", vertical="center")
            larguras[j] = max(larguras[j], len(str(col)) + 2)
        linha += 1
        for _, row in df.iterrows():
            for j, col in enumerate(df.columns, start=1):
                val = row[col]
                if isinstance(val, (list, tuple)):
                    val = str(val)
                elif isinstance(val, float):
                    val = None if math.isnan(val) else round(val, 4)
                c = ws.cell(row=linha, column=j, value=val)
                c.font, c.border = body_font, borda
                larguras[j] = max(larguras[j], len(str(val)) + 2)
            linha += 1
        linha += 2
    for j, w in larguras.items():
        ws.column_dimensions[get_column_letter(j)].width = min(max(w, 10), 40)

    destino = os.path.join(pasta, "comparacao_final.xlsx")
    wb.save(destino)
    print(f"\n  Exportado: {destino}")


def main():
    p = argparse.ArgumentParser(description="Avaliação do Benchmark RAG (guid)")
    p.add_argument("--ficheiro")
    p.add_argument("--pasta")
    p.add_argument("--comparar", action="store_true")
    p.add_argument("--so-retrieval", action="store_true")
    p.add_argument("--langsmith", action="store_true")
    args = p.parse_args()

    if args.comparar:
        comparar(args.pasta)
    elif args.ficheiro:
        avaliar_ficheiro(args.ficheiro, args.so_retrieval, args.langsmith)
    elif args.pasta:
        fs = [f for f in glob.glob(os.path.join(args.pasta, "*.json"))
              if "_avaliacao" not in f and "comparacao" not in f]
        for f in sorted(fs):
            avaliar_ficheiro(f, args.so_retrieval, args.langsmith)
        comparar(args.pasta)
    else:
        p.print_help()


if __name__ == "__main__":
    main()